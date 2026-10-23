from pathlib import Path


def parse_uploaded_content(path: Path, suffix: str) -> dict:
    """Extract a lightweight structure summary before content-specific normalization."""
    if suffix in {".xlsx", ".xls"}:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        sheets = []
        for worksheet in workbook.worksheets:
            values = list(worksheet.iter_rows(values_only=True))
            keys = [str(value).strip() if value is not None else "" for value in (values[0] if values else ())]
            labels = [str(value).strip() if value is not None else "" for value in (values[1] if len(values) > 1 else ())]
            has_template_labels = bool(keys and labels and len(values) > 1)
            data_rows = values[2:] if has_template_labels else values[1:]
            records = []
            if worksheet.title != "字段说明" and keys and any(keys):
                for row in data_rows:
                    record = {key: value for key, value in zip(keys, row) if key and value not in (None, "")}
                    if record:
                        records.append(record)
            sheets.append({
                "name": worksheet.title,
                "rows": worksheet.max_row or 0,
                "columns": worksheet.max_column or 0,
                "headers": labels or keys,
                "data_rows": len(records),
                "records_preview": records[:5],
            })
        workbook.close()
        return {"format": "excel", "sheets": sheets, "sheet_count": len(sheets)}

    if suffix == ".docx":
        from docx import Document

        document = Document(path)
        paragraphs = [item.text.strip() for item in document.paragraphs if item.text.strip()]
        return {
            "format": "word",
            "paragraph_count": len(paragraphs),
            "table_count": len(document.tables),
            "text_preview": "\n".join(paragraphs[:8])[:2000],
        }

    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
        return {
            "format": "pdf",
            "page_count": len(reader.pages),
            "character_count": len(text),
            "text_preview": text[:2000],
        }

    return {"format": suffix.lstrip("."), "status": "requires_legacy_converter"}
